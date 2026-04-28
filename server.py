#!/usr/bin/env python3
"""
Sistema de Extracción de Facturas + Registro PTR Mall del Sol
- API Key ingresada por el usuario (sin keys hardcodeadas)
- Modelo principal: gemini-2.5-flash
- Procesa múltiples PDFs de forma SECUENCIAL (1 por 1 internamente)
- Búsqueda de SOLPED/Requisición en CUALQUIER parte del PDF
- Módulo PTR: extrae datos de PTR y genera Ficha Técnica de Autorizaciones
- Clasificación automática por proveedor (Infraestructura / Técnico)
"""

from flask import Flask, request, jsonify, send_file, render_template, session
from flask_cors import CORS
import fitz  # PyMuPDF
import json
import re
import io
import os
import time
import traceback
from datetime import datetime, date
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import secrets

try:
    from google import genai
    from google.genai import types
    GEMINI_OK = True
except ImportError:
    GEMINI_OK = False

app = Flask(__name__, static_folder='static')
# Usa SECRET_KEY de env var para que las sesiones sobrevivan reinicios en Render
# Si no hay variable configurada, genera una aleatoria (solo para desarrollo)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
CORS(app, supports_credentials=True)

# ─────────────────────────────────────────────────────────
# CONFIGURACIÓN DE MODELOS
# Modelo principal: gemini-2.5-flash
# Fallback: gemini-2.0-flash → gemini-2.0-flash-lite
# ─────────────────────────────────────────────────────────
CONFIG = {
    "model_primary": "gemini-2.5-flash",
    "models": ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-2.0-flash-lite"],
    "max_retries": 3,
    "retry_delay": 20,
}


# ─────────────────────────────────────────────────────────
# MANEJO DE API KEY DE SESIÓN
# ─────────────────────────────────────────────────────────
def get_session_key():
    """Obtiene la API key de la sesión actual."""
    return session.get("api_key", "")

def set_session_key(key):
    """Guarda la API key en la sesión."""
    session["api_key"] = key.strip()


# ─────────────────────────────────────────────────────────
# TABLA DE CUENTAS → RUBRO
# ─────────────────────────────────────────────────────────
TABLA_CUENTAS = {
    "52-1-08-01-01": "BAÑOS",
    "52-1-08-01-02": "RIEGO",
    "52-1-08-01-03": "FUENTES",
    "52-1-08-01-04": "EQUIPOS DE SENTINA",
    "52-1-08-01-05": "EQUIPOS CONTRA INCENDIO",
    "52-1-08-01-06": "INSTALACIONES ELECTRICAS",
    "52-1-08-01-07": "TRANSFORMADORES Y SUBESTACIONES",
    "52-1-08-01-08": "EQUIPOS DE RADIO",
    "52-1-08-01-09": "VEHICULOS SKYJACK",
    "52-1-08-01-10": "CISTERNAS",
    "52-1-08-01-11": "SISTEMA DE AGUA POTABLE",
    "52-1-08-01-12": "AIRE ACONDICIONADO",
    "52-1-08-01-13": "ESCALERAS Y ELEVADORES",
    "52-1-08-01-14": "HERRAMIENTAS",
    "52-1-08-01-15": "AUDIO",
    "52-1-08-01-16": "CENTRAL TELEFONICA",
    "52-1-08-01-17": "CAMARAS DE SEGURIDAD",
    "52-1-08-01-18": "GENERACION",
    "52-1-08-01-19": "UPS CENTRAL",
    "52-1-08-01-20": "PUERTAS ELECTRICAS",
    "52-1-08-01-21": "MEDIDORES DE AGUA",
    "52-1-08-01-22": "VALLAS",
    "52-1-08-01-23": "VIDEO PATIO DE COMIDA",
    "52-1-08-01-24": "LOCALES VACIOS Y NUEVAS INSTALACIONES",
    "52-1-08-01-25": "SISTEMA DE GAS",
    "52-1-08-01-26": "VOZ Y DATOS",
    "52-1-08-01-27": "IMPREVISTOS",
    "52-1-24-02-01": "FOCOS Y LAMPARAS",
    "52-1-08-02-04": "AGUA POTABLE",
    "41-2-01-02-06": "ANTENAS",
    "41-2-01-02-07": "MANTENIMIENTO AC",
    "43-5-01-01-05": "REEMBOLSO DE GASTOS",
    "52-1-18-01-07": "TASA RECOLECCION DE BASURA",
    "52-1-28-01-01": "SUMINISTROS DE OFICINA",
    "52-1-18-01-04": "INTERAGUA",
    "10-2-01-03-01": "CAPEX (O ACTIVOS EN PROCESO)",
    "10-2-02-05-01": "CONSTRUCCIONES EN PROCESO",
    "52-2-04-01-05": "VARIOS",
    "52-4-02-01-07": "GASTOS INTERMEDIARIOS",
    "52-1-18-01-03": "PLANILLAS",
    "52-1-28-03-13": "OTROS GASTOS ADMINISTRATIVOS",
    "10-1-03-04-01": "DIESEL/INVENTARIO INSUMOS",
    "52-1-28-07-09": "PROTECTORES MRN",
    "52-2-11-03-03": "IMPRESIONES PUBLICIDAD CLIENTES",
    "52-2-12-01-09": "GASTOS PANORAMIX",
    "52-1-28-02-05": "LIMPIEZA AMPLIACION",
    "52-1-08-01-29": "MANTENIMIENTOS TECNICOS AMPLIACION",
    "52-1-08-02-32": "MANTENIMIENTOS OPERATIVOS AMPLIACION",
    "52-1-24-02-11": "INSUMOS AMPLIACION",
    "52-1-18-01-09": "ENERGIA AMPLIACION",
    "52-1-18-01-10": "AGUA POTABLE AMPLIACION",
    "52-2-12-01-01": "ACTIVACIONES EXPERIENCIA",
    "52-1-28-02-01": "TECNICO",
}

TABLA_TEXTO = "\n".join(f"  {k} → {v}" for k, v in TABLA_CUENTAS.items())

# ─────────────────────────────────────────────────────────
# TABLAS DE PROVEEDORES PTR
# ─────────────────────────────────────────────────────────
PROVEEDORES_INFRAESTRUCTURA = [
    "AGUIRRE ZAVALA CRISTIAN",
    "BRAUCOS",
    "CONSCIVIM SA",
    "CUATRICOMIA",
    "FERNANDEZ PABLO A. COUMENGES",
    "HEROSA",
    "ING. ERNESTO SALTOS",
    "JAVIER DIEZ COMUNICACIÓN VISUAL CIA. LTDA.",
    "MANLIM S.A.",
    "MASSUH SUAREZ JOSE GABRIEL",
    "MEDINA ARCENTALES JUAN CARLOS",
    "MEJIA YAGUAL JUSTINO",
    "METALMURI",
    "MUROCOMS S.A.S.",
    "NATSOLUTIONS",
    "OPSA",
    "ORLANDO ANDRE VALVERDE CASTRO (HEROSA)",
    "PUNTOTRADE SAS",
    "RENOVO HOGAR S.A.S",
    "VALVERDE CASTRO ORLANDO ANDRE (HEROSA)",
]

PROVEEDORES_TECNICO = [
    "AGUILA SECURITY/URBAPARK",
    "BRUGUESA",
    "CONSDIASACORP S.A.",
    "CSR INGENIERIA",
    "DOLDER",
    "GPS",
    "INDUSUR",
    "JAVIER DIEZ COMUNICACIÓN VISUAL CIA. LTDA.",
    "MENDOTEL",
    "METALTHUNDER S.A.",
    "S3T TELQUALITY",
    "SERGELIMP",
    "TOTEM",
]

# Tipos de PTR según número de acceso
TIPOS_PTR = {
    "N/A": "NO APLICA ACCESO A CUBIERTA- NO APLICA EPP",
    "1": "PTR: USO DE EPP (CASCO, CHALECO Y BOTAS)",
    "2": "PTR: ACCESO TERRAZA / PERRERA",
    "3": "PTR: 4TO PISO - INSPECCIÓN CON JEFE TÉCNICO",
    "4": "PTR OTROS (IZAJE) USO DE EPP",
    "5": "PTR CALIENTE- USO DE EPP",
    "6": "PTR CONFINADO- USO DE EPP",
    "7": "PTR ELÉCTRICO/ENERGIZADO- USO DE EPP",
    "8": "PTR ALTURA- USO DE EPP",
    "9": "PTR EXCAVACIONES",
}

def classify_provider(nombre):
    """Clasifica un proveedor en INFRAESTRUCTURA o TÉCNICO."""
    nombre_upper = nombre.upper().strip()
    for p in PROVEEDORES_INFRAESTRUCTURA:
        if p.upper() in nombre_upper or nombre_upper in p.upper():
            return "INFRAESTRUCTURA"
    for p in PROVEEDORES_TECNICO:
        if p.upper() in nombre_upper or nombre_upper in p.upper():
            return "TÉCNICO"
    # Búsqueda parcial
    for p in PROVEEDORES_INFRAESTRUCTURA:
        words = [w for w in p.upper().split() if len(w) > 3]
        if any(w in nombre_upper for w in words):
            return "INFRAESTRUCTURA"
    for p in PROVEEDORES_TECNICO:
        words = [w for w in p.upper().split() if len(w) > 3]
        if any(w in nombre_upper for w in words):
            return "TÉCNICO"
    return "INFRAESTRUCTURA"  # default

def get_week_number(fecha_str):
    """Calcula el número de semana del año para una fecha DD/MM/YYYY."""
    try:
        if isinstance(fecha_str, str):
            parts = fecha_str.replace("-", "/").split("/")
            if len(parts) == 3:
                if len(parts[2]) == 4:  # DD/MM/YYYY
                    d = datetime(int(parts[2]), int(parts[1]), int(parts[0]))
                else:  # YYYY/MM/DD
                    d = datetime(int(parts[0]), int(parts[1]), int(parts[2]))
        else:
            d = fecha_str
        return d.isocalendar()[1]
    except:
        return ""


# ─────────────────────────────────────────────────────────
# PROMPT FACTURA PRINCIPAL
# ─────────────────────────────────────────────────────────
def build_prompt_factura():
    return f"""Eres un experto en extracción de datos de facturas comerciales ecuatorianas.
Se te entregan TODAS LAS PÁGINAS del documento en orden (factura en pág.1, orden de compra en pág.2-3, etc.).

TAREA: Extrae los campos y devuelve ÚNICAMENTE un objeto JSON válido (sin markdown, sin ```, sin texto extra).

{{
  "fecha": "Fecha de EMISIÓN de la factura. Formato DD/MM/AAAA. Busca 'Fecha Emisión' o 'Fecha y hora de Autorización'. Solo la fecha de la factura, no de la orden.",
  "proveedor": "Nombre COMERCIAL de quien EMITE la factura (tiene su RUC y logo arriba a la izquierda). NUNCA pongas 'MOBILSOL', 'INMOBILIARIA DEL SOL' ni variantes — esas somos NOSOTROS el comprador.",
  "factura": "Número completo de factura. Formato: 001-001-000002702.",
  "solped": "SOLPED o REQUISICIÓN. BUSCA EN ABSOLUTAMENTE TODAS LAS PÁGINAS sin excepción. Patrón: SOL seguido de 5+ dígitos (ej: SOL0012163). Busca en: campo 'REQUISICIÓN', 'OBSERVACIONES', 'Información Adicional', pie de página, notas, cualquier campo. Si NO encuentras → ''",
  "orden": "Número de Orden de Compra. BUSCA EN TODAS LAS PÁGINAS. Patrón: OC seguido de 5+ dígitos (ej: OC0015541). Si no existe → ''",
  "rubro": "Elige el rubro MÁS APROPIADO según qué se compra. Devuelve SOLO el nombre exacto de esta tabla:\n{TABLA_TEXTO}",
  "cuenta": "Código contable que corresponde al rubro elegido (ej: 52-1-08-01-12). Si no estás seguro → ''",
  "proyecto": "SOLO si el documento dice explícitamente 'PROYECTO:' o 'Proyecto:' seguido de nombre. Si no aparece → ''",
  "descripcion": "Descripción del trabajo/material. Busca primero en OBSERVACIONES de la OC, luego en ítems de la factura. Elige la más completa.",
  "sub_total": "Monto subtotal sin IVA. Solo número decimal. Ej: 27.53",
  "iva": "Monto del IVA. Solo número decimal. Ej: 4.13. Si es 0 → 0.00",
  "total": "Total de la factura. Solo número decimal. Ej: 31.66"
}}

REGLAS CRÍTICAS:
1. PROVEEDOR = quien factura (parte superior de la factura con su RUC). MOBILSOL = comprador → NUNCA proveedor.
2. SOLPED: busca en TODAS las páginas sin excepción. En OC suele estar en campo "REQUISICIÓN". También puede aparecer como "SOL-XXXXXXX" con guión.
3. ORDEN OC: busca en TODAS las páginas. Puede estar en la factura o en la OC.
4. Ejemplos de RUBRO: ventilador/AC → AIRE ACONDICIONADO (52-1-08-01-12); focos → FOCOS Y LAMPARAS (52-1-24-02-01); generador → GENERACION (52-1-08-01-18).
5. PROYECTO: vacío si no dice literalmente "Proyecto:" en el documento.
6. Datos no encontrados → vacío "". NUNCA inventes valores.
7. Responde SOLO el JSON, sin ningún texto adicional."""


# ─────────────────────────────────────────────────────────
# PROMPT PTR
# ─────────────────────────────────────────────────────────
def build_prompt_ptr():
    tipos_str = "\n".join(f"  {k}: {v}" for k, v in TIPOS_PTR.items())
    proveedores_infra = "\n".join(f"  - {p}" for p in PROVEEDORES_INFRAESTRUCTURA)
    proveedores_tec = "\n".join(f"  - {p}" for p in PROVEEDORES_TECNICO)

    return f"""Eres un experto en extracción de datos de documentos PTR (Permiso de Trabajo en Altura/Riesgo) del Mall del Sol en Ecuador.
Se te entregan TODAS LAS PÁGINAS del documento PTR.

TAREA: Extrae los campos y devuelve ÚNICAMENTE un objeto JSON válido (sin markdown, sin ```, sin texto extra).

{{
  "nombre_comercial": "Nombre del contratista/empresa que realiza el trabajo. Busca en 'Trabajo realizado por:', 'Nombre del Contratista:', 'Contratista:'. NUNCA pongas 'Inmobiliaria del Sol', 'Mall del Sol' ni 'MDS' — esos son el dueño. En mayúsculas.",
  "fecha_inicio": "Fecha de inicio del trabajo. Busca 'Fecha de emisión' o 'Fecha:' en la segunda hoja. Formato DD/MM/YYYY.",
  "fecha_fin": "Fecha límite/validez del permiso. Busca 'Permiso válido hasta'. Si es la misma que inicio, repite. Formato DD/MM/YYYY.",
  "descripcion_trabajo": "Descripción completa del trabajo. Busca en 'DETALLE DEL TRABAJO A EJECUTAR' (segunda hoja) que es la más completa. Si no está, usa 'Descripción del trabajo:' de la primera página.",
  "tipo_ptr_numero": "Número del tipo de PTR marcado con X en la segunda hoja. Opciones: N/A, 1, 2, 3, 4, 5, 6, 7, 8, 9. Busca cuál está marcado con X en 'TIPO DE PERMISO DE TRABAJO SOLICITADO'. Si no hay marca clara, infiere según descripción.",
  "personal": "Lista de TODO el personal autorizado con nombre completo y número de cédula/identificación. Formato: 'APELLIDO NOMBRE: XXXXXXXXXX'. Busca en 'DETALLE DEL PERSONAL CONTRATISTA AUTORIZADO'. Incluye TODOS sin excepción.",
  "area_trabajo": "Área o ubicación donde se realiza el trabajo. Busca 'Área de Trabajo:', 'Ubicación del trabajo:', 'Fachada', 'Local', etc.",
  "responsable_supervisor": "Nombre del responsable/supervisor del contratista. Busca 'Nombre de Responsable / Supervisor:' o 'Supervisor:'.",
  "contacto_celular": "Número de celular de contacto. Busca 'Celular de Contacto:'."
}}

TIPOS DE PTR para referencia:
{tipos_str}

PROVEEDORES INFRAESTRUCTURA (si el contratista coincide con alguno, categoría = INFRAESTRUCTURA):
{proveedores_infra}

PROVEEDORES TÉCNICO:
{proveedores_tec}

REGLAS CRÍTICAS:
1. nombre_comercial: quien REALIZA el trabajo (contratista), NO Inmobiliaria del Sol/Mall del Sol/MDS.
2. fecha_inicio y fecha_fin: suelen ser iguales en un PTR de un solo día. Busca bien en ambas páginas.
3. descripcion_trabajo: prioriza la de la segunda hoja ('DETALLE DEL TRABAJO A EJECUTAR') que es más detallada.
4. personal: incluye ABSOLUTAMENTE TODOS los trabajadores listados con nombre y cédula.
5. tipo_ptr_numero: busca la X marcada en la segunda hoja en 'TIPO DE PERMISO DE TRABAJO SOLICITADO'.
6. Responde SOLO el JSON, sin texto adicional."""


# ─────────────────────────────────────────────────────────
# FUNCIONES CORE DE PROCESAMIENTO
# ─────────────────────────────────────────────────────────

def pdf_to_images(pdf_bytes, zoom=2.0):
    """Convierte todas las páginas del PDF a JPEG."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    images = []
    mat = fitz.Matrix(zoom, zoom)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=mat)
        jpg = pix.tobytes("jpeg")
        images.append(jpg)
        print(f"    Pág {i+1}: {len(jpg)//1024} KB")
    n = len(doc)
    doc.close()
    return images, n


def clean_json(raw):
    """Extrae JSON puro de la respuesta."""
    text = raw.strip()
    text = re.sub(r'^```(?:json)?\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'\s*```\s*$', '', text, flags=re.MULTILINE)
    start = text.find('{')
    end   = text.rfind('}')
    if start != -1 and end != -1:
        text = text[start:end+1]
    return text.strip()


def normalize_factura(data):
    """Normaliza y valida los datos de factura extraídos."""
    required = ["fecha", "proveedor", "factura", "solped", "orden",
                "rubro", "cuenta", "proyecto", "descripcion", "sub_total", "iva", "total"]
    for f in required:
        if f not in data or data[f] is None:
            data[f] = ""
        else:
            data[f] = str(data[f]).strip()

    # Bloquear MOBILSOL como proveedor
    prov = data.get("proveedor", "").lower()
    if any(x in prov for x in ["mobilsol", "movilsol", "inmobiliaria del sol"]):
        data["proveedor"] = ""

    # Normalizar SOLPED
    s = data.get("solped", "")
    if s and not re.match(r'^SOL\d+$', s, re.IGNORECASE):
        m = re.search(r'SOL[-]?\d{5,10}', s, re.IGNORECASE)
        data["solped"] = m.group(0).replace("-", "").upper() if m else ""
    elif not s:
        for campo in ["descripcion", "proyecto", "orden"]:
            texto = data.get(campo, "")
            m = re.search(r'SOL[-]?\d{5,10}', texto, re.IGNORECASE)
            if m:
                data["solped"] = m.group(0).replace("-", "").upper()
                break

    # Normalizar ORDEN
    o = data.get("orden", "")
    if o and not re.match(r'^OC\d+$', o, re.IGNORECASE):
        m = re.search(r'OC[-]?\d{5,10}', o, re.IGNORECASE)
        data["orden"] = m.group(0).replace("-", "").upper() if m else ""

    # Sincronizar RUBRO ↔ CUENTA
    rubro = data.get("rubro", "").upper().strip()
    cuenta = data.get("cuenta", "").strip()
    if cuenta and cuenta in TABLA_CUENTAS:
        data["rubro"] = TABLA_CUENTAS[cuenta]
    elif rubro:
        for c, d in TABLA_CUENTAS.items():
            if d.upper() == rubro:
                data["cuenta"] = c
                data["rubro"] = d
                break

    # Normalizar montos
    for f in ["sub_total", "iva", "total"]:
        val = data.get(f, "").replace("$", "").replace(",", "").strip()
        try:
            float(val)
        except (ValueError, TypeError):
            val = ""
        data[f] = val

    return data


def normalize_ptr(data):
    """Normaliza datos extraídos de un PTR."""
    required = ["nombre_comercial", "fecha_inicio", "fecha_fin", "descripcion_trabajo",
                "tipo_ptr_numero", "personal", "area_trabajo", "responsable_supervisor", "contacto_celular"]
    for f in required:
        if f not in data or data[f] is None:
            data[f] = ""
        elif isinstance(data[f], list):
            data[f] = "\n".join(str(x).strip().strip("'\"") for x in data[f] if x)
        else:
            data[f] = str(data[f]).strip()

    # Limpiar campo personal
    personal = data.get("personal", "")
    if personal.startswith("[") and personal.endswith("]"):
        try:
            import ast
            parsed = ast.literal_eval(personal)
            if isinstance(parsed, list):
                personal = "\n".join(str(x).strip() for x in parsed if x)
        except Exception:
            personal = personal.strip("[]")
            personal = re.sub(r"['\"]", "", personal)
            personal = re.sub(r",\s*", "\n", personal)
    data["personal"] = personal.strip()

    # Categoría y semana
    nombre = data.get("nombre_comercial", "")
    data["categoria"] = classify_provider(nombre)
    data["semana"] = str(get_week_number(data.get("fecha_inicio", "")))

    # Prefijo PTR en descripción
    desc = data.get("descripcion_trabajo", "")
    if desc and not desc.upper().startswith("PTR"):
        data["descripcion_trabajo"] = f"PTR: {desc}"

    return data


def extract_with_gemini(pdf_bytes, prompt_fn, api_key, doc_type="factura"):
    """
    Extrae datos usando Gemini con la API key proporcionada por el usuario.
    Modelo principal: gemini-2.5-flash
    Fallback: gemini-2.0-flash → gemini-2.0-flash-lite
    """
    if not GEMINI_OK:
        return None, "Librería google-genai no instalada. Ejecute: pip install google-genai"

    if not api_key:
        return None, "No hay API key configurada. Ingresa tu API key de Google Gemini en la barra superior."

    prompt = prompt_fn()

    print(f"  Convirtiendo PDF a imágenes...")
    images, num_pages = pdf_to_images(pdf_bytes)
    if not images:
        return None, "No se pudieron extraer imágenes del PDF"

    print(f"  {num_pages} página(s) extraídas")

    base_parts = [types.Part.from_bytes(data=img, mime_type="image/jpeg") for img in images]
    base_parts.append(types.Part.from_text(text=prompt))

    last_error = "Error desconocido"
    models_tried = []

    for model in CONFIG["models"]:
        models_tried.append(model)
        max_retries = CONFIG["max_retries"]

        for attempt in range(1, max_retries + 1):
            print(f"  [{model}] Intento {attempt}/{max_retries}...")

            try:
                client = genai.Client(api_key=api_key)
                t0 = time.time()
                response = client.models.generate_content(
                    model=model,
                    contents=base_parts,
                    config=types.GenerateContentConfig(temperature=0.1)
                )
                elapsed = time.time() - t0
                raw = response.text.strip()
                print(f"  [{model}] OK en {elapsed:.1f}s | {len(raw)} chars")

                clean = clean_json(raw)
                data = json.loads(clean)

                if doc_type == "ptr":
                    data = normalize_ptr(data)
                else:
                    data = normalize_factura(data)

                data["_paginas"] = num_pages
                data["_model"] = model
                return data, None

            except json.JSONDecodeError as e:
                last_error = f"Error parseando JSON de Gemini: {e}"
                print(f"  [{model}] JSON error: {e}")
                break

            except Exception as e:
                err_str = str(e)
                last_error = err_str

                if "RESOURCE_EXHAUSTED" in err_str or "429" in err_str:
                    print(f"  [{model}] Cuota agotada para esta key")
                    if attempt >= max_retries:
                        break
                    time.sleep(CONFIG["retry_delay"])
                    continue

                elif "PERMISSION_DENIED" in err_str or "403" in err_str or "API_KEY_INVALID" in err_str:
                    return None, "API Key inválida o sin permisos. Verifica que sea correcta en Google AI Studio."

                elif "NOT_FOUND" in err_str or "404" in err_str:
                    print(f"  [{model}] Modelo no disponible, probando siguiente...")
                    break

                else:
                    print(f"  [{model}] Error: {err_str[:200]}")
                    if attempt >= max_retries:
                        break

    return None, (
        f"No se pudo procesar el documento. "
        f"Modelos probados: {', '.join(models_tried)}. "
        f"Último error: {last_error[:200]}. "
        f"Verifica tu API key o espera unos minutos si agotaste la cuota."
    )


# ─────────────────────────────────────────────────────────
# GENERACIÓN DE EXCEL — FACTURAS
# ─────────────────────────────────────────────────────────
def create_excel_facturas(registros):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control Facturas"

    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    hdr_font  = Font(color="FFFFFF", bold=True, size=10)
    alt_fill  = PatternFill("solid", fgColor="EBF3FB")
    norm_fill = PatternFill("solid", fgColor="FFFFFF")
    warn_fill = PatternFill("solid", fgColor="FFFBEB")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )

    headers = ["FECHA","PROVEEDOR","FACTURA","SOLPED","ORDEN",
               "RUBRO","CUENTA","PROYECTO","DESCRIPCIÓN DEL TRABAJO",
               "SUB TOTAL","IVA","TOTAL FACTURAR"]
    widths  = [12, 30, 22, 14, 14, 30, 20, 28, 55, 12, 10, 14]
    campos  = ["fecha","proveedor","factura","solped","orden",
               "rubro","cuenta","proyecto","descripcion",
               "sub_total","iva","total"]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill      = hdr_fill
        c.font      = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = thin
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 35

    for ri, reg in enumerate(registros, 2):
        fill = alt_fill if ri % 2 == 0 else norm_fill
        for ci, campo in enumerate(campos, 1):
            val = reg.get(campo, "")
            c   = ws.cell(row=ri, column=ci, value=val)
            if not val and campo in ("fecha", "proveedor", "factura", "solped"):
                c.fill = warn_fill
            else:
                c.fill = fill
            c.border    = thin
            c.alignment = Alignment(vertical="center", wrap_text=(ci == 9))
        ws.row_dimensions[ri].height = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ─────────────────────────────────────────────────────────
# GENERACIÓN DE EXCEL — PTR / FICHA TÉCNICA
# ─────────────────────────────────────────────────────────
def create_excel_ptr(registros):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ficha Técnica PTR"

    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    hdr_font  = Font(color="FFFFFF", bold=True, size=10)
    alt_fill  = PatternFill("solid", fgColor="EBF3FB")
    norm_fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin")
    )

    title_fill = PatternFill("solid", fgColor="1F4E79")
    ws.merge_cells("A1:M1")
    t = ws.cell(row=1, column=1, value="FICHA TÉCNICA AUTORIZACIONES COORDINACIÓN GENERAL MALL DEL SOL")
    t.fill = title_fill
    t.font = Font(color="FFFFFF", bold=True, size=12)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    headers = [
        "NOMBRE COMERCIAL", "CATEGORIA- ÁREA", "SEMANA",
        "FECHA- INICIO", "FECHA- FINAL", "HORA ACCESO", "HORA FIN/ACCESO",
        "DESCRIPCIÓN DEL TRABAJO A REALIZAR:", "ACCESOS-SIA",
        "NOMBRE DE EMPRESA", "PERSONAL AUTORIZADO CON ACCESO:",
        "FECHA APROBADO", "RESPONSABLE:"
    ]
    widths = [25, 18, 9, 14, 14, 12, 12, 50, 12, 25, 45, 14, 18]
    campos = [
        "nombre_comercial", "categoria", "semana",
        "fecha_inicio", "fecha_fin", "hora_acceso", "hora_fin",
        "descripcion_trabajo", "accesos_sia",
        "nombre_empresa", "personal",
        "fecha_aprobado", "responsable"
    ]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=2, column=col, value=h)
        c.fill      = hdr_fill
        c.font      = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = thin
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[2].height = 35

    for ri, reg in enumerate(registros, 3):
        fill = alt_fill if ri % 2 == 1 else norm_fill
        for ci, campo in enumerate(campos, 1):
            val = reg.get(campo, "")
            c   = ws.cell(row=ri, column=ci, value=val)
            c.fill   = fill
            c.border = thin
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[ri].height = 60

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ─────────────────────────────────────────────────────────
# RUTAS API — CONFIGURACIÓN DE API KEY
# ─────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/ptr")
def ptr_page():
    return render_template("ptr.html")


@app.route("/api/set-key", methods=["POST"])
def set_api_key():
    """Guarda la API key del usuario en la sesión."""
    body = request.get_json() or {}
    key = body.get("api_key", "").strip()
    if not key:
        return jsonify({"error": "API key vacía"}), 400
    if not key.startswith("AIza"):
        return jsonify({"error": "Formato de API key inválido. Debe comenzar con 'AIza'"}), 400
    set_session_key(key)
    print(f"[KEY] API key configurada: {key[:8]}...")
    return jsonify({
        "success": True,
        "message": "API key guardada correctamente",
        "preview": key[:8] + "..." + key[-4:],
        "model": CONFIG["model_primary"],
    })


@app.route("/api/get-key-status", methods=["GET"])
def get_key_status():
    """Retorna el estado de la API key de sesión (sin revelar la key completa)."""
    key = get_session_key()
    has_key = bool(key)
    return jsonify({
        "has_key": has_key,
        "preview": (key[:8] + "..." + key[-4:]) if has_key else None,
        "model": CONFIG["model_primary"],
        "models_fallback": CONFIG["models"],
    })


@app.route("/api/health")
def health():
    key = get_session_key()
    return jsonify({
        "status": "ok",
        "model": CONFIG["model_primary"],
        "models": CONFIG["models"],
        "gemini_ready": GEMINI_OK,
        "api_key_set": bool(key),
        "cuentas": len(TABLA_CUENTAS),
    })


@app.route("/api/config", methods=["GET"])
def get_config():
    key = get_session_key()
    return jsonify({
        "model": CONFIG["model_primary"],
        "models": CONFIG["models"],
        "api_key_set": bool(key),
        "cuentas": len(TABLA_CUENTAS),
    })


@app.route("/api/cuentas")
def get_cuentas():
    return jsonify(TABLA_CUENTAS)


@app.route("/api/proveedores")
def get_proveedores():
    return jsonify({
        "infraestructura": PROVEEDORES_INFRAESTRUCTURA,
        "tecnico": PROVEEDORES_TECNICO,
    })


# ─────────────────────────────────────────────────────────
# RUTAS API — EXTRACCIÓN
# ─────────────────────────────────────────────────────────

@app.route("/api/extract", methods=["POST"])
def extract_factura():
    """Extrae datos de factura usando la API key del usuario."""
    try:
        if "file" not in request.files:
            return jsonify({"error": "No se recibió ningún archivo"}), 400

        file = request.files["file"]
        if not file.filename:
            return jsonify({"error": "Nombre de archivo vacío"}), 400
        if not file.filename.lower().endswith(".pdf"):
            return jsonify({"error": "Solo se aceptan archivos PDF"}), 400

        pdf_bytes = file.read()
        if not pdf_bytes:
            return jsonify({"error": "Archivo PDF vacío"}), 400

        api_key = get_session_key()
        if not api_key:
            return jsonify({"error": "⚠ No hay API key configurada. Ingresa tu API key de Gemini en la barra superior."}), 400

        print(f"\n{'='*60}")
        print(f"[FACTURA] {file.filename} ({len(pdf_bytes)//1024} KB)")

        data, error = extract_with_gemini(pdf_bytes, build_prompt_factura, api_key, "factura")
        if error:
            return jsonify({"error": error}), 500

        print(f"[OK] proveedor={data.get('proveedor')} | solped={data.get('solped')}")
        return jsonify({"success": True, "data": data})

    except Exception as e:
        print(f"[ERROR] {traceback.format_exc()}")
        return jsonify({"error": str(e)}), 500


@app.route("/api/extract-ptr", methods=["POST"])
def extract_ptr():
    """Extrae datos de un PTR usando la API key del usuario."""
    try:
        if "file" not in request.files:
            return jsonify({"error": "No se recibió ningún archivo"}), 400

        file = request.files["file"]
        if not file.filename:
            return jsonify({"error": "Nombre de archivo vacío"}), 400
        if not file.filename.lower().endswith(".pdf"):
            return jsonify({"error": "Solo se aceptan archivos PDF"}), 400

        pdf_bytes = file.read()
        if not pdf_bytes:
            return jsonify({"error": "Archivo PDF vacío"}), 400

        api_key = get_session_key()
        if not api_key:
            return jsonify({"error": "⚠ No hay API key configurada. Ingresa tu API key de Gemini en la barra superior."}), 400

        print(f"\n{'='*60}")
        print(f"[PTR] {file.filename} ({len(pdf_bytes)//1024} KB)")

        data, error = extract_with_gemini(pdf_bytes, build_prompt_ptr, api_key, "ptr")
        if error:
            return jsonify({"error": error}), 500

        data["nombre_empresa"]  = data.get("nombre_comercial", "")
        data["hora_acceso"]     = "VARIOS"
        data["hora_fin"]        = "VARIOS"
        data["accesos_sia"]     = ""
        data["fecha_aprobado"]  = datetime.now().strftime("%d/%m/%Y")
        data["responsable"]     = data.get("categoria", "INFRAESTRUCTURA")

        print(f"[OK] contratista={data.get('nombre_comercial')} | cat={data.get('categoria')}")
        return jsonify({"success": True, "data": data})

    except Exception as e:
        print(f"[ERROR] {traceback.format_exc()}")
        return jsonify({"error": str(e)}), 500


@app.route("/api/export", methods=["POST"])
def export_excel():
    """Exporta registros de facturas a Excel."""
    try:
        body = request.get_json()
        registros = body.get("registros", [])
        if not registros:
            return jsonify({"error": "Sin registros para exportar"}), 400

        buf   = create_excel_facturas(registros)
        fname = f"Control_Facturas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True, download_name=fname)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/export-ptr", methods=["POST"])
def export_excel_ptr():
    """Exporta registros PTR a Excel (Ficha Técnica)."""
    try:
        body = request.get_json()
        registros = body.get("registros", [])
        if not registros:
            return jsonify({"error": "Sin registros para exportar"}), 400

        buf   = create_excel_ptr(registros)
        fname = f"Ficha_Tecnica_PTR_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True, download_name=fname)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    # Render inyecta el puerto via $PORT; en local usamos 3000
    port = int(os.environ.get("PORT", 3000))
    print("=" * 60)
    print("  MOBILSOL — Control Facturas + PTR Manager")
    print("=" * 60)
    print(f"  Modelo principal : {CONFIG['model_primary']}")
    print(f"  Modelos fallback : {' → '.join(CONFIG['models'][1:])}")
    print(f"  Cuentas          : {len(TABLA_CUENTAS)} rubros")
    print(f"  Proveedores Infra: {len(PROVEEDORES_INFRAESTRUCTURA)}")
    print(f"  Proveedores Tec  : {len(PROVEEDORES_TECNICO)}")
    print(f"  API Key          : Ingresada por cada usuario en la UI")
    print(f"  Puerto           : {port}")
    print("=" * 60)
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
